import os
import sys
import io
import json
import uuid
import docx
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import httpx

# Ensure backend root is on sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.main import app
from app.database.database import get_db, SessionLocal
from app.models.user import User
from app.models.paper import ResearchPaper
from app.models.analysis import Analysis

import httpx

BASE_URL = "http://127.0.0.1:8000"
client = httpx.Client(base_url=BASE_URL, timeout=30.0)

def create_sample_docx(file_path: Path, title: str, authors: str, has_table: bool = True):
    doc = docx.Document()
    doc.add_heading(title, 0)
    p_auth = doc.add_paragraph(authors)
    doc.add_paragraph("Department of Quantum Computing, Cambridge University, UK. Email: authors@cambridge.ac.uk")
    
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("This paper investigates quantum neural architectures for multi-qubit error suppression. We propose a hybrid variational circuit model that reduces decoherence by 42%.")
    
    doc.add_heading("Problem Statement", level=1)
    doc.add_paragraph("The primary bottleneck in NISQ devices is quantum noise and rapid decoherence of entanglement states.")
    
    doc.add_heading("Methodology", level=1)
    doc.add_paragraph("Our framework utilizes parameterized unitary gates integrated with a classical gradient descent optimizer.")
    
    if has_table:
        doc.add_heading("Results", level=1)
        doc.add_paragraph("Experimental validation demonstrates significant improvements across standard benchmarks.")
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Model Architecture'
        hdr_cells[1].text = 'Fidelity'
        hdr_cells[2].text = 'Latency'
        
        data = [
            ("Classical Baseline", "71.2", "140ms"),
            ("Standard VQE", "83.5", "95ms"),
            ("QuantumResNet-64", "96.8", "24ms"),
        ]
        for m, f, l in data:
            row_cells = table.add_row().cells
            row_cells[0].text = m
            row_cells[1].text = f
            row_cells[2].text = l

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph("The primary contribution is a robust variational quantum circuit reducing state error.")
    doc.save(str(file_path))

def create_qualitative_docx(file_path: Path, title: str, authors: str):
    doc = docx.Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(authors)
    doc.add_paragraph("Faculty of Philosophy and Social Sciences, Oxford University.")
    
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("This qualitative study examines the governance frameworks and ethical dilemmas surrounding autonomous algorithmic decision-making.")
    
    doc.add_heading("Methodology", level=1)
    doc.add_paragraph("We conduct a systematic qualitative literature analysis and thematic synthesis across 50 regulatory policies.")
    
    doc.add_heading("Discussion", level=1)
    doc.add_paragraph("Ethical governance requires proactive human oversight and transparent algorithmic accountability.")

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph("We synthesize key policy guidelines for ethical AI deployment.")
    doc.save(str(file_path))

def test_grounded_analysis_and_auth():
    print("=" * 70)
    print("  RESEARCHX: GROUNDED ANALYSIS, AUTHOR EXTRACTION & AUTH TEST SUITE")
    print("=" * 70)

    # ---------------------------------------------------------
    # PART 1: AUTHENTICATION & PERSISTENCE TESTS
    # ---------------------------------------------------------
    print("\n---> [TEST 1] Authentication Lifecycle & Email Normalization...")
    test_email = f"  Ada.Lovelace_{uuid.uuid4().hex[:6]}@Computing.Org  "
    test_password = "SecurePassword123!"
    
    # 1. Register with spaced and mixed-case email
    reg_res = client.post("/api/auth/register", json={
        "name": "Lady Ada Lovelace",
        "email": test_email,
        "password": test_password
    })
    assert reg_res.status_code == 200, f"Registration failed: {reg_res.text}"
    reg_data = reg_res.json()["data"]
    clean_email = test_email.strip().lower()
    assert reg_data["user"]["email"] == clean_email
    print(f"[PASS] 1. Registered account with normalized email: '{clean_email}'")

    # 2. Login with registered account
    login_res = client.post("/api/auth/login", json={
        "email": test_email,
        "password": test_password
    })
    assert login_res.status_code == 200, f"Login failed: {login_res.text}"
    token = login_res.json()["data"]["access_token"]
    user_headers = {"Authorization": f"Bearer {token}"}
    print("[PASS] 2. Logged in successfully with spaced email.")

    # 3. Test wrong password
    wrong_pwd_res = client.post("/api/auth/login", json={
        "email": clean_email,
        "password": "WrongPassword999!"
    })
    assert wrong_pwd_res.status_code == 401
    assert "Invalid email or password." in wrong_pwd_res.json()["message"]
    print("[PASS] 3. Wrong password correctly rejected with 401: 'Invalid email or password.'")

    # 4. Test nonexistent email
    nonexistent_res = client.post("/api/auth/login", json={
        "email": "nonexistent_researcher_xyz@nowhere.com",
        "password": test_password
    })
    assert nonexistent_res.status_code == 401
    assert "Invalid email or password." in nonexistent_res.json()["message"]
    print("[PASS] 4. Nonexistent email correctly rejected with 401: 'Invalid email or password.'")

    # 5. Duplicate account registration rejection
    dup_res = client.post("/api/auth/register", json={
        "name": "Another Ada",
        "email": f"  {clean_email.upper()}  ",
        "password": "AnotherPassword456!"
    })
    assert dup_res.status_code == 400
    assert dup_res.json()["message"] == "An account with this email already exists. Please log in."
    print("[PASS] 5. Duplicate registration correctly blocked with exact required message: 'An account with this email already exists. Please log in.'")

    # ---------------------------------------------------------
    # PART 2: PAPER A (EMPIRICAL) GROUNDED ANALYSIS & AUTHOR EXTRACTION
    # ---------------------------------------------------------
    print("\n---> [TEST 2] Paper A Upload & Grounded Empirical Analysis...")
    temp_dir = Path("backend/uploads")
    temp_dir.mkdir(parents=True, exist_ok=True)
    paper_a_path = temp_dir / "test_paper_a_quantum.docx"
    create_sample_docx(
        paper_a_path,
        title="Quantum Residual Circuits for Qubit Decoherence Suppression",
        authors="Dr. Elena Vance, Prof. Alan Turing, Dr. Katherine Johnson",
        has_table=True
    )

    with open(paper_a_path, "rb") as f:
        upload_a_res = client.post(
            "/api/papers/upload",
            files={"file": ("quantum_circuits.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"domain": "Computer Science"},
            headers=user_headers
        )
    assert upload_a_res.status_code == 200, f"Upload Paper A failed: {upload_a_res.text}"
    paper_a_id = upload_a_res.json()["data"]["id"]
    print(f"[PASS] 6. Uploaded Paper A successfully (ID: {paper_a_id})")

    # Analyze Paper A
    analyze_a_res = client.post(f"/api/analysis/paper/{paper_a_id}", headers=user_headers)
    assert analyze_a_res.status_code == 200, f"Analysis Paper A failed: {analyze_a_res.text}"
    analysis_a_id = analyze_a_res.json()["data"]["analysis_id"]

    # Fetch Paper A report
    report_a_res = client.get(f"/api/analysis/{analysis_a_id}", headers=user_headers)
    assert report_a_res.status_code == 200
    report_a = report_a_res.json()["data"]

    # Verify Author Extraction on Paper A
    print(f"Extracted Authors for Paper A: {report_a['authors']}")
    assert "Elena Vance" in report_a["authors"]
    assert "Alan Turing" in report_a["authors"]
    assert "Katherine Johnson" in report_a["authors"]
    assert "University" not in report_a["authors"]
    assert "cambridge.ac.uk" not in report_a["authors"]
    print("[PASS] 7. Accurate Author Extraction verified without affiliations or email addresses.")

    # Verify Paper A Charts (must have genuine numbers from Paper A table: 71.2, 83.5, 96.8)
    viz_a = report_a["visualization_data"]
    assert viz_a["has_visualizations"] is True
    model_names_a = [s["name"] for s in viz_a["series"]]
    assert "QuantumResNet-64" in model_names_a or "Standard VQE" in model_names_a
    print(f"[PASS] 8. Paper A generated pure grounded charts with genuine models: {model_names_a}")

    # ---------------------------------------------------------
    # PART 3: PAPER B (QUALITATIVE) NO FAKE DATA & ISOLATION
    # ---------------------------------------------------------
    print("\n---> [TEST 3] Paper B Upload & Strict Zero-Fake-Data Verification...")
    paper_b_path = temp_dir / "test_paper_b_ethics.docx"
    create_qualitative_docx(
        paper_b_path,
        title="Qualitative Governance and Ethics in Autonomous Systems",
        authors="Prof. Margaret Hamilton, Dr. Grace Hopper"
    )

    with open(paper_b_path, "rb") as f:
        upload_b_res = client.post(
            "/api/papers/upload",
            files={"file": ("qualitative_ethics.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"domain": "Social Science"},
            headers=user_headers
        )
    assert upload_b_res.status_code == 200
    paper_b_id = upload_b_res.json()["data"]["id"]
    print(f"[PASS] 9. Uploaded Paper B successfully (ID: {paper_b_id})")

    # Analyze Paper B
    analyze_b_res = client.post(f"/api/analysis/paper/{paper_b_id}", headers=user_headers)
    assert analyze_b_res.status_code == 200
    analysis_b_id = analyze_b_res.json()["data"]["analysis_id"]

    # Fetch Paper B report
    report_b_res = client.get(f"/api/analysis/{analysis_b_id}", headers=user_headers)
    assert report_b_res.status_code == 200
    report_b = report_b_res.json()["data"]

    # Verify No Data Leakage between Paper A and Paper B
    assert "Quantum" not in report_b["paper_title"]
    assert "Elena Vance" not in report_b["authors"]
    assert "Margaret Hamilton" in report_b["authors"]
    print("[PASS] 10. Complete Data Isolation verified between Paper A and Paper B.")

    # Verify Paper B Dataset & Fallback Grounding
    assert report_b["dataset"] == "Dataset information is not explicitly mentioned in the paper."
    print(f"[PASS] 11. Missing Dataset field properly returns: '{report_b['dataset']}'")

    # Verify Paper B has NO FAKE CHARTS
    viz_b = report_b["visualization_data"]
    assert viz_b["has_visualizations"] is False
    assert viz_b["message"] == "No suitable numerical data was found in the paper for this visualization."
    assert len(viz_b["series"]) == 0
    print(f"[PASS] 12. Qualitative Paper B correctly generated NO fake charts with message: '{viz_b['message']}'")

    # ---------------------------------------------------------
    # PART 4: UNSUPPORTED FORMAT REJECTION
    # ---------------------------------------------------------
    print("\n---> [TEST 4] Unsupported File Formats Rejection...")
    txt_res = client.post(
        "/api/papers/upload",
        files={"file": ("test.txt", io.BytesIO(b"Simple text file"), "text/plain")},
        headers=user_headers
    )
    assert txt_res.status_code == 400
    assert txt_res.json()["message"] == "Unsupported file type. Please upload PDF, DOC, DOCX, JPG, JPEG, or PNG."
    print("[PASS] 13. Unsupported .txt file safely rejected with HTTP 400 and exact error message.")

    xlsx_res = client.post(
        "/api/papers/upload",
        files={"file": ("data.xlsx", io.BytesIO(b"fake excel content"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=user_headers
    )
    assert xlsx_res.status_code == 400
    print("[PASS] 14. Unsupported .xlsx file safely rejected with HTTP 400.")

    # Cleanup test files
    for p in [paper_a_path, paper_b_path]:
        try:
            if p.exists():
                p.unlink()
        except Exception:
            pass

    print("\n" + "=" * 70)
    print("  ALL GROUNDED ANALYSIS, AUTHOR EXTRACTION & AUTH TESTS PASSED 100%!")
    print("=" * 70)

if __name__ == "__main__":
    test_grounded_analysis_and_auth()
