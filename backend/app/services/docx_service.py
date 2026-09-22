import re
import zipfile
from pathlib import Path
from typing import List, Optional
import docx
from app.services.pdf_service import PDFService

class DocxService:
    @staticmethod
    def extract_text_from_docx(file_path: Path) -> str:
        """
        Extract clean, structured plain text from a Microsoft Word .docx file.
        Preserves paragraph order, extracts headings, handles tables,
        and excludes empty lines and binary artifacts.
        """
        if not file_path.exists():
            raise FileNotFoundError(f"Document file not found at: {file_path}")

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise ValueError(f"Failed to open .docx file. The document may be corrupted or encrypted: {str(e)}")

        content_blocks: List[str] = []

        # 1. Extract paragraphs in order
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                # If paragraph style is a heading, demarcate clearly
                style_name = (para.style.name if para.style else "").lower()
                if "heading" in style_name or "title" in style_name:
                    content_blocks.append(f"\n{text}\n")
                else:
                    content_blocks.append(text)

        # 2. Extract structured table data
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                # Deduplicate adjacent merged cells in the same row
                seen_cell_texts = []
                for cell in row.cells:
                    ctext = cell.text.strip()
                    if ctext and (not seen_cell_texts or ctext != seen_cell_texts[-1]):
                        seen_cell_texts.append(ctext)
                if seen_cell_texts:
                    table_rows.append(" | ".join(seen_cell_texts))
            if table_rows:
                content_blocks.append("\n" + "\n".join(table_rows) + "\n")

        full_text = "\n\n".join(content_blocks).strip()

        if not full_text:
            raise ValueError("The uploaded .docx document does not contain any readable text.")

        return PDFService.clean_text(full_text)

    @staticmethod
    def extract_text_from_doc(file_path: Path) -> str:
        """
        Extract readable plain text from a legacy Microsoft Word .doc file.
        Uses pure-Python stream and binary parsing without requiring external OS packages.
        """
        if not file_path.exists():
            raise FileNotFoundError(f"Document file not found at: {file_path}")

        # Check if file is actually a renamed .docx (ZIP format)
        try:
            if zipfile.is_zipfile(str(file_path)):
                return DocxService.extract_text_from_docx(file_path)
        except Exception:
            pass

        # Read binary content for OLE2 binary Word documents
        try:
            with open(file_path, "rb") as f:
                content = f.read()
        except Exception as e:
            raise ValueError(f"Could not read .doc file: {str(e)}")

        # Extract readable UTF-16LE and ASCII text blocks
        extracted_pieces = []
        
        # 1. UTF-16LE chunks
        utf16_matches = re.findall(b'(?:[\x20-\x7e\r\n\t]\x00){4,}', content)
        for match in utf16_matches:
            try:
                decoded = match.decode('utf-16le', errors='ignore').strip()
                if len(decoded) > 10 and any(c.isalpha() for c in decoded):
                    extracted_pieces.append(decoded)
            except Exception:
                pass

        # 2. ASCII/Latin-1 strings
        ascii_matches = re.findall(b'[\x20-\x7e\r\n\t]{6,}', content)
        for match in ascii_matches:
            try:
                decoded = match.decode('latin-1', errors='ignore').strip()
                # Exclude common OLE metadata strings
                if len(decoded) > 15 and not decoded.startswith(('WordDocument', 'SummaryInformation', 'DocumentSummaryInformation', 'Microsoft Word', 'Normal.dot')):
                    if any(c.isalpha() for c in decoded):
                        extracted_pieces.append(decoded)
            except Exception:
                pass

        # Deduplicate and assemble
        seen = set()
        clean_pieces = []
        for p in extracted_pieces:
            normalized = p.strip()
            if normalized and normalized not in seen and len(normalized) > 5:
                seen.add(normalized)
                clean_pieces.append(normalized)

        if not clean_pieces:
            raise ValueError("The legacy .doc file could not be parsed. Please convert or save the file in .docx format.")

        full_text = "\n\n".join(clean_pieces)
        return PDFService.clean_text(full_text)
